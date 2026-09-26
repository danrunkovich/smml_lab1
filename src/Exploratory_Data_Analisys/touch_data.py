import json
import os
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import seaborn as sns
from sklearn.preprocessing import MinMaxScaler, OneHotEncoder


def main() -> None:
    json_path = "../../dataset/audi_a6_dataset_fuel_type_RLEncoded.json"
    if not os.path.exists(json_path):
        print(f"ОШИБКА: Файл {json_path} не найден!")
        return

    print(f"Загрузка данных из: {json_path}")
    with open(json_path, "r", encoding="utf-8") as file:
        raw_data = json.load(file)

    df = pd.DataFrame(raw_data)

    for col in ["drive", "transmission", "fuel_type", "body_type"]:
        df[col] = df[col].astype(str).str.strip().str.lower()

    num_cols = ["year", "mileage_km", "power_hp", "engine_vol_l", "price_rub"]
    scaler = MinMaxScaler(feature_range=(0, 1))

    scaled_numeric = scaler.fit_transform(df[num_cols])
    scaled_num_df = pd.DataFrame(
        scaled_numeric,
        columns=[
            "scaled_year",
            "scaled_mileage",
            "scaled_power",
            "scaled_engine_vol",
            "scaled_price",
        ],
    )

    categories_spec = [
        ["полный", "передний"],
        ["вариатор", "автомат", "механика", "робот"],
        ["дизель", "бензин", "другое"],
        ["универсал", "седан"],
    ]
    cat_cols = ["drive", "transmission", "fuel_type", "body_type"]

    ohe = OneHotEncoder(
        categories=categories_spec,
        drop="first",
        sparse_output=False,
        handle_unknown="ignore",
    )
    ohe_matrix = ohe.fit_transform(df[cat_cols])
    active_ohe_columns = ohe.get_feature_names_out(cat_cols)

    ohe_active_df = pd.DataFrame(ohe_matrix, columns=active_ohe_columns, dtype=int)
    active_transformed_df = pd.concat([scaled_num_df, ohe_active_df], axis=1)

    json_output_df = pd.concat([df[["url", "price_rub"]], active_transformed_df], axis=1)

    output_json_path = "../../dataset/audi_a6_dataset_transformed.json"
    os.makedirs(os.path.dirname(output_json_path), exist_ok=True)
    json_output_df.to_json(output_json_path, orient="records", force_ascii=False, indent=2)
    print(f"Новый JSON (только активные n - 1 фичи) сохранен: {output_json_path}")


    os.makedirs("./plots/transformed_data_plots", exist_ok=True)
    sns.set_theme(style="whitegrid")

    corr_matrix = active_transformed_df.corr()

    #sorted_features = corr_matrix["scaled_price"].sort_values(ascending=False).index
    #sorted_corr_matrix = corr_matrix.loc[sorted_features, sorted_features]

    plt.figure(figsize=(12, 10))
    sns.heatmap(
        corr_matrix,
        cmap="coolwarm",
        annot=True,
        fmt=".2f",
        linewidths=0.5,
        cbar_kws={"shrink": 0.8},
    )
    plt.title(
        "Матрица парных корреляций",
        fontsize=14,
        fontweight="bold",
    )
    plt.tight_layout()

    heatmap_path = ("./plots/transformed_data_plots/transformed_correlation_heatmap.png")
    plt.savefig(heatmap_path, dpi=300)
    plt.close()
    print(f"Тепловая карта успешно сохранена: {heatmap_path}")



    excel_cat_df = pd.DataFrame(index=df.index)

    for col_name, cats in zip(cat_cols, categories_spec):
        base_cat = cats[0]
        # Столбец базовой категории
        excel_cat_df[f"{col_name}_{base_cat}"] = 0
        # Активные столбцы (n - 1)
        for cat in cats[1:]:
            col_name_active = f"{col_name}_{cat}"
            excel_cat_df[col_name_active] = ohe_active_df[col_name_active]

    excel_export_df = pd.concat([scaled_num_df, excel_cat_df], axis=1)

    excel_path = "../../dataset/Audi_A6_dataset.xlsx"
    if not os.path.exists(excel_path):
        excel_path = "../dataset/Audi_A6_dataset.xlsx"

    if os.path.exists(excel_path):
        print(f"Обновление Excel файла: {excel_path}")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        START_COL = 16

        for c_idx, col_name in enumerate(excel_export_df.columns, start=START_COL):
            ws.cell(row=1, column=c_idx, value=col_name)

        for r_idx, row_values in enumerate(excel_export_df.values, start=2):
            for c_idx, val in enumerate(row_values, start=START_COL):
                cell_val = (
                    round(float(val), 4) if isinstance(val, float) else int(val)
                )
                ws.cell(row=r_idx, column=c_idx, value=cell_val)

        wb.save(excel_path)
        print(f"Excel успешно обновлен: начиная с 16-го столбца записано {len(excel_export_df.columns)} колонок")
    else:
        print(f"Внимание: Excel файл {excel_path} не найден. Проверьте путь к файлу")

    print("\nвсе гуд пушка ракета торпеда")


if __name__ == "__main__":
    main()