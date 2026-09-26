import json
import os
import openpyxl


def main() -> None:
    # 1. Поиск исходного JSON файла
    json_path = "../../dataset/audi_a6_dataset_input.json"
    if not os.path.exists(json_path):
        print(f"ОШИБКА: Файл {json_path} не найден!")
        return

    print(f"Чтение исходных данных из: {json_path}")
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    rare_categories = {"электро", "гибрид"}
    replaced_count = 0

    for item in data:
        fuel = str(item.get("fuel_type", "")).strip().lower()
        if fuel in rare_categories:
            item["fuel_type"] = "другое"
            replaced_count += 1
        else:
            item["fuel_type"] = fuel

    print(f"Успешно обработано: {replaced_count} редких значений заменено на 'другое'.")

    output_json_path = "../../dataset/audi_a6_dataset_fuel_type_RLEncoded.json"
    os.makedirs(os.path.dirname(output_json_path), exist_ok=True)

    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Новый JSON сохранен: {output_json_path}")

    excel_path = "../../dataset/Audi_A6_dataset.xlsx"

    if os.path.exists(excel_path):
        print(f"Обновление Excel файла: {excel_path}")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        COL_IDX = 13
        ws.cell(row=1, column=COL_IDX, value="fuel_type_preprocessed")

        for row_idx, item in enumerate(data, start=2):
            ws.cell(row=row_idx, column=COL_IDX, value=item["fuel_type"])

        wb.save(excel_path)
        print(f"Готово! В столбец {COL_IDX} добавлена предобработанная фича fuel_type.")
    else:
        print(f"Внимание: Excel файл {excel_path} не найден, обновлен только JSON.")


if __name__ == "__main__":
    main()